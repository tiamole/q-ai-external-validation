from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PACKAGE_ROOT = Path(__file__).resolve().parents[1]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_configuration(package_root: Path) -> dict[str, Any]:
    config_path = package_root / "config" / "frozen-analysis.yml"
    with config_path.open(encoding="utf-8") as source:
        return yaml.safe_load(source)


def expected_source_hash(package_root: Path, relative_path: str) -> str:
    manifest_path = package_root / "data-manifest" / "source-files.csv"
    with manifest_path.open(newline="", encoding="utf-8") as source:
        rows = list(csv.DictReader(source))
    matches = [row for row in rows if row["local_path"] == relative_path]
    if len(matches) != 1:
        raise ValueError(f"Expected one manifest record for {relative_path!r}, found {len(matches)}")
    return matches[0]["local_sha256"].lower()


def integer_score(value: Any, participant_id: str, statement_id: str) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Non-numeric score at {participant_id}/{statement_id}: {value!r}")
    if not float(value).is_integer():
        raise ValueError(f"Non-integer score at {participant_id}/{statement_id}: {value!r}")
    return int(value)


def normalize_sort(
    participant_id: str,
    row_number: int,
    statement_ids: list[str],
    raw_values: tuple[Any, ...],
    expected_distribution: Counter[int],
    repair_config: dict[str, Any],
) -> tuple[list[int], dict[str, Any] | None]:
    scores = [
        integer_score(value, participant_id, statement_id)
        for statement_id, value in zip(statement_ids, raw_values, strict=True)
    ]
    legal_scores = set(expected_distribution)
    invalid_indexes = [index for index, score in enumerate(scores) if score not in legal_scores]
    repair: dict[str, Any] | None = None

    if invalid_indexes:
        if len(invalid_indexes) != 1:
            raise ValueError(f"Multiple out-of-range scores for {participant_id}: {invalid_indexes}")
        index = invalid_indexes[0]
        statement_id = statement_ids[index]
        source_cell = f"{get_column_letter(index + 2)}{row_number}"
        expected_cell = repair_config["source_cell"].split("!")[-1]
        expected_anomaly = (
            participant_id == repair_config["participant_id"]
            and statement_id == repair_config["statement_id"]
            and scores[index] == repair_config["raw_value"]
            and source_cell == expected_cell
        )
        if not expected_anomaly:
            raise ValueError(
                f"Unprespecified out-of-range score at {participant_id}/{statement_id} "
                f"({source_cell}): {scores[index]}"
            )
        raw_value = scores[index]
        scores[index] = int(repair_config["normalized_value"])
        repair = {
            "source_cell": repair_config["source_cell"],
            "participant_id": participant_id,
            "statement_id": statement_id,
            "raw_value": raw_value,
            "normalized_value": scores[index],
            "rule": "unique forced-distribution completion",
        }

    observed_distribution = Counter(scores)
    if observed_distribution != expected_distribution:
        raise ValueError(
            f"Forced-distribution mismatch for {participant_id}: "
            f"expected {dict(sorted(expected_distribution.items()))}, "
            f"observed {dict(sorted(observed_distribution.items()))}"
        )
    return scores, repair


def write_matrix(path: Path, statement_ids: list[str], rows: list[tuple[str, list[int]]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as destination:
        writer = csv.writer(destination, lineterminator="\n")
        writer.writerow(["participant", *statement_ids])
        for participant_id, scores in rows:
            writer.writerow([participant_id, *scores])


def validate_and_normalize(
    package_root: Path = PACKAGE_ROOT,
    output_root: Path | None = None,
) -> dict[str, Any]:
    config = load_configuration(package_root)
    input_config = config["input"]
    relative_workbook = input_config["workbook"]
    workbook_path = package_root / relative_workbook
    expected_hash = expected_source_hash(package_root, relative_workbook)
    actual_hash = sha256(workbook_path)
    if actual_hash != expected_hash:
        raise ValueError(
            f"Source workbook SHA-256 mismatch: expected {expected_hash}, observed {actual_hash}"
        )

    statement_config = input_config["statement_columns"]
    statement_ids = [
        f"{statement_config['prefix']}{number}"
        for number in range(statement_config["first"], statement_config["last"] + 1)
    ]
    expected_distribution = Counter(
        {int(item["score"]): int(item["count"]) for item in config["forced_distribution"]}
    )
    participant_pattern = re.compile(input_config["participant_id_pattern"])

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet_name = input_config["worksheet"]
        if worksheet_name not in workbook.sheetnames:
            raise ValueError(f"Required worksheet not found with exact name {worksheet_name!r}")
        worksheet = workbook[worksheet_name]
        headers = [worksheet.cell(1, column).value for column in range(1, len(statement_ids) + 2)]
        expected_headers = [input_config["participant_column"], *statement_ids]
        if headers != expected_headers:
            raise ValueError(f"Unexpected Q-sort headers: {headers!r}")

        source_rows: list[tuple[int, str, tuple[Any, ...]]] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=len(expected_headers), values_only=True), start=2
        ):
            participant_value = values[0]
            if isinstance(participant_value, str) and participant_pattern.fullmatch(
                participant_value.strip()
            ) and any(value is not None for value in values[1:]):
                source_rows.append((row_number, participant_value.strip(), values[1:]))
    finally:
        workbook.close()

    expected_participants = input_config["expected_participants"]
    if len(source_rows) != expected_participants:
        raise ValueError(f"Expected {expected_participants} participant rows, found {len(source_rows)}")
    participant_ids = [participant_id for _, participant_id, _ in source_rows]
    if len(participant_ids) != len(set(participant_ids)):
        raise ValueError("Participant identifiers are not unique")

    normalized_rows: list[tuple[str, list[int]]] = []
    repairs: list[dict[str, Any]] = []
    for row_number, participant_id, raw_values in source_rows:
        scores, repair = normalize_sort(
            participant_id,
            row_number,
            statement_ids,
            raw_values,
            expected_distribution,
            config["prespecified_repair"],
        )
        normalized_rows.append((participant_id, scores))
        if repair:
            repairs.append(repair)

    if len(repairs) != 1:
        raise ValueError(f"Expected exactly one prespecified repair, applied {len(repairs)}")

    duplicate_groups: list[list[str]] = []
    sorts_by_signature: dict[tuple[int, ...], list[str]] = {}
    for participant_id, scores in normalized_rows:
        sorts_by_signature.setdefault(tuple(scores), []).append(participant_id)
    for participant_group in sorts_by_signature.values():
        if len(participant_group) > 1:
            duplicate_groups.append(participant_group)

    output_base = output_root or package_root
    normalized_dir = output_base / "data" / "normalized" / "pidgeon-2025"
    integrity_dir = output_base / "outputs" / "integrity"
    normalized_path = normalized_dir / "qsorts_repaired.csv"
    sensitivity_path = normalized_dir / "qsorts_without_P47.csv"
    repair_log_path = integrity_dir / "repair-log.csv"
    report_path = integrity_dir / "pidgeon-input-validation.json"

    write_matrix(normalized_path, statement_ids, normalized_rows)
    sensitivity_rows = [row for row in normalized_rows if row[0] != "P47"]
    write_matrix(sensitivity_path, statement_ids, sensitivity_rows)

    integrity_dir.mkdir(parents=True, exist_ok=True)
    with repair_log_path.open("w", newline="", encoding="utf-8") as destination:
        writer = csv.DictWriter(destination, fieldnames=list(repairs[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(repairs)

    report = {
        "schema_version": "1.0",
        "status": "pass",
        "study_id": config["study"]["id"],
        "input": {
            "path": relative_workbook,
            "sha256": actual_hash,
            "worksheet": input_config["worksheet"],
        },
        "observed": {
            "participants": len(normalized_rows),
            "statements": len(statement_ids),
            "repairs": repairs,
            "duplicate_sort_groups": duplicate_groups,
        },
        "derived_files": {
            "repaired_matrix": str(normalized_path.relative_to(output_base)).replace("\\", "/"),
            "exclusion_matrix": str(sensitivity_path.relative_to(output_base)).replace("\\", "/"),
            "repair_log": str(repair_log_path.relative_to(output_base)).replace("\\", "/"),
        },
    }
    with report_path.open("w", encoding="utf-8", newline="\n") as destination:
        json.dump(report, destination, indent=2, sort_keys=True)
        destination.write("\n")
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate and normalize the Pidgeon Q-sort matrix.")
    parser.add_argument("--package-root", type=Path, default=PACKAGE_ROOT)
    args = parser.parse_args()
    report = validate_and_normalize(args.package_root.resolve())
    print(
        f"Validation {report['status']}: {report['observed']['participants']} participants, "
        f"{report['observed']['statements']} statements, "
        f"{len(report['observed']['repairs'])} logged repair."
    )


if __name__ == "__main__":
    main()